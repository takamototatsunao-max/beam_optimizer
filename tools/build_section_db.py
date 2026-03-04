#!/usr/bin/env python3
"""
Build SECTION_DB sheet from extracted CSV, based on an input template workbook.
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _to_float(v: str):
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    return float(s.replace(",", ""))


def _read_rows(csv_path: Path) -> List[Dict[str, object]]:
    out: List[Dict[str, object]] = []
    with csv_path.open("r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            out.append({
                "SectionName": str(row.get("SectionName", "")).strip(),
                "Series": str(row.get("Series", "")).strip(),
                "SourcePage": int(float(row.get("SourcePage", "0") or 0)),
                "Nominal": str(row.get("Nominal", "")).strip(),
                "h_mm": _to_float(row.get("h_mm")),
                "b_mm": _to_float(row.get("b_mm")),
                "tw_mm": _to_float(row.get("tw_mm")),
                "tf_mm": _to_float(row.get("tf_mm")),
                "r_mm": _to_float(row.get("r_mm")),
                "A_cm2": _to_float(row.get("A_cm2")),
                "Ix_cm4": _to_float(row.get("Ix_cm4")),
                "Zx_cm3": _to_float(row.get("Zx_cm3")),
                "UnitMass_kg_m": _to_float(row.get("UnitMass_kg_m")),
                "w_g_kN_m": _to_float(row.get("w_g_kN_m")),
            })
    out = [r for r in out if r["SectionName"]]
    out.sort(key=lambda r: (float(r["w_g_kN_m"]), str(r["SectionName"])))
    for i, r in enumerate(out, start=1):
        r["Rank"] = i
    return out


def _set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, help="template workbook path (e.g. input_b17.xlsx)")
    ap.add_argument("--csv", required=True, help="section csv path from extract_h_sections_pdf.py")
    ap.add_argument("--out", required=True, help="output workbook path (e.g. input_b18.xlsx)")
    ap.add_argument(
        "--blank-wg-count",
        type=int,
        default=0,
        help="number of top-ranked rows where w_g is intentionally left blank (default: 0)",
    )
    args = ap.parse_args()

    template_path = Path(args.template)
    csv_path = Path(args.csv)
    out_path = Path(args.out)
    blank_wg_count = max(0, int(args.blank_wg_count))

    if not template_path.exists():
        raise SystemExit(f"Template workbook not found: {template_path}")
    if not csv_path.exists():
        raise SystemExit(f"CSV not found: {csv_path}")

    rows = _read_rows(csv_path)
    if not rows:
        raise SystemExit("No section rows loaded from CSV.")

    wb = load_workbook(template_path)
    if "SECTION_DB" in wb.sheetnames:
        del wb["SECTION_DB"]
    ws = wb.create_sheet("SECTION_DB")

    headers = [
        "Use",
        "Rank",
        "SectionName",
        "Series",
        "SourcePage",
        "Nominal",
        "h [mm]",
        "b [mm]",
        "tw [mm]",
        "tf [mm]",
        "r [mm]",
        "A [cm2]",
        "Ix [cm4]",
        "Zx [cm3]",
        "UnitMass [kg/m]",
        "w_g [kN/m]",
        "Av [mm2]",
        "Note",
    ]
    ws.append(headers)
    _set_col_widths(ws, [8, 8, 22, 10, 10, 12, 10, 10, 10, 10, 8, 10, 12, 10, 14, 12, 12, 32])

    blanked = 0
    for row in rows:
        h = float(row["h_mm"])
        tw = float(row["tw_mm"])
        tf = float(row["tf_mm"])
        av_mm2 = tw * (h - 2.0 * tf)
        wg = float(row["w_g_kN_m"])
        if blanked < blank_wg_count:
            wg_cell = None
            blanked += 1
        else:
            wg_cell = wg
        ws.append([
            True,
            int(row["Rank"]),
            row["SectionName"],
            row["Series"],
            int(row["SourcePage"]),
            row["Nominal"],
            float(row["h_mm"]),
            float(row["b_mm"]),
            float(row["tw_mm"]),
            float(row["tf_mm"]),
            float(row["r_mm"]) if row["r_mm"] is not None else None,
            float(row["A_cm2"]) if row["A_cm2"] is not None else None,
            float(row["Ix_cm4"]) if row["Ix_cm4"] is not None else None,
            float(row["Zx_cm3"]) if row["Zx_cm3"] is not None else None,
            float(row["UnitMass_kg_m"]) if row["UnitMass_kg_m"] is not None else None,
            wg_cell,
            av_mm2,
            f"from binran_chapter05 p{int(row['SourcePage'])}",
        ])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"rows={len(rows)} blank_wg={blanked} written={out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
